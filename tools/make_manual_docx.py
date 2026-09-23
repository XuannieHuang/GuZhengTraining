# -*- coding: utf-8 -*-
"""把 操作手冊.md 轉成 Word 檔 → 操作手冊.docx

給不看 GitHub 的人用。手冊或截圖更新後重跑一次即可：

    py tools/make_manual_docx.py

只支援這份手冊實際用到的 Markdown 語法：標題、段落、粗體、行內程式碼、
連結、引言、表格、圖片（<img src width>）、程式碼區塊、分隔線、項目符號。
"""
import os, re, sys, io

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, '操作手冊.md')
DST = os.path.join(ROOT, '操作手冊.docx')

CJK = '微軟正黑體'
MONO = 'Consolas'
IMG_W = Inches(2.9)          # 手機截圖放這個寬度剛好，一頁放得下說明


def set_font(run, name=CJK, size=None, bold=None, color=None, mono=False):
    f = name if not mono else MONO
    run.font.name = f
    run._element.rPr.rFonts.set(qn('w:eastAsia'), CJK)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def shade(cell, hex_color):
    el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear')
    el.set(qn('w:fill'), hex_color)
    cell._tc.get_or_add_tcPr().append(el)


INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')


def add_inline(par, text, size=10.5, base_bold=False):
    """處理粗體、行內程式碼、連結（連結只保留文字＋網址）"""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            set_font(par.add_run(part[2:-2]), size=size, bold=True)
        elif part.startswith('`') and part.endswith('`'):
            set_font(par.add_run(part[1:-1]), size=size - 0.5, mono=True,
                     color=(0xB0, 0x2A, 0x37))
        elif part.startswith('['):
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', part)
            set_font(par.add_run(m.group(1)), size=size, bold=base_bold,
                     color=(0x1A, 0x5F, 0xB4))
            if not m.group(2).startswith('#'):
                set_font(par.add_run(f'（{m.group(2)}）'), size=size - 1.5,
                         color=(0x70, 0x70, 0x70))
        else:
            set_font(par.add_run(part), size=size, bold=base_bold)


def split_row(line):
    return [c.strip() for c in line.strip().strip('|').split('|')]


def main():
    lines = io.open(SRC, encoding='utf-8').read().split('\n')
    doc = Document()

    st = doc.styles['Normal']
    st.font.name = CJK
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), CJK)
    for s in ('Heading 1', 'Heading 2', 'Heading 3', 'Heading 4'):
        doc.styles[s].element.rPr.rFonts.set(qn('w:eastAsia'), CJK)
        doc.styles[s].font.color.rgb = RGBColor(0x1F, 0x22, 0x28)

    i, n_img, n_tbl = 0, 0, 0
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()

        if not s or s == '---':
            i += 1
            continue

        # 圖片
        m = re.match(r'<img src="([^"]+)"(?:\s+width="(\d+)")?>', s)
        if m:
            path = os.path.join(ROOT, m.group(1).replace('/', os.sep))
            if os.path.exists(path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(path, width=IMG_W)
                n_img += 1
            else:
                print(f'  ⚠️ 找不到圖片：{m.group(1)}')
            i += 1
            continue

        # 標題
        m = re.match(r'^(#{1,4})\s+(.*)', s)
        if m:
            lvl, txt = len(m.group(1)), m.group(2)
            h = doc.add_heading('', level=lvl)
            add_inline(h, txt, size={1: 20, 2: 15, 3: 12.5, 4: 11}[lvl], base_bold=True)
            i += 1
            continue

        # 程式碼區塊
        if s.startswith('```'):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                buf.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(8)
            set_font(p.add_run('\n'.join(buf)), size=9.5, mono=True)
            continue

        # 表格
        if s.startswith('|') and i + 1 < len(lines) and re.match(r'^\|[\s:|-]+\|$', lines[i + 1].strip()):
            head = split_row(s)
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append(split_row(lines[i]))
                i += 1
            t = doc.add_table(rows=1, cols=len(head))
            t.style = 'Table Grid'
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for c, txt in zip(t.rows[0].cells, head):
                c.text = ''
                add_inline(c.paragraphs[0], txt, size=10, base_bold=True)
                shade(c, 'EEF0F2')
            for r in rows:
                cells = t.add_row().cells
                for c, txt in zip(cells, r + [''] * (len(head) - len(r))):
                    c.text = ''
                    add_inline(c.paragraphs[0], txt, size=10)
            doc.add_paragraph()
            n_tbl += 1
            continue

        # 引言
        if s.startswith('>'):
            body = []
            while i < len(lines) and lines[i].strip().startswith('>'):
                body.append(lines[i].strip().lstrip('>').strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_inline(p, ' '.join(x for x in body if x), size=10)
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x55, 0x5A, 0x62)
            continue

        # 項目符號 / 編號
        m = re.match(r'^([-*]|\d+\.)\s+(.*)', s)
        if m:
            p = doc.add_paragraph(style='List Bullet' if m.group(1) in '-*' else 'List Number')
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, m.group(2))
            i += 1
            continue

        # 一般段落
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        add_inline(p, s)
        i += 1

    doc.save(DST)
    print(f'✅ {DST}')
    print(f'   {n_img} 張圖、{n_tbl} 個表格，{os.path.getsize(DST)/1024/1024:.1f} MB')


if __name__ == '__main__':
    main()
